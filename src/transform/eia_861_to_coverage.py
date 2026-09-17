"""EIA-861 thin coverage path (layer G) — utility/state + best-effort utility→county.

Reads extracted Service_Territory_2024.xlsx and Utility_Data_2024.xlsx under
data/raw/eia/. Matches EIA county *names* to Census county FIPS via gazetteer
name normalization. Never invents FIPS for ambiguous or unmatched names.

This is **utility service-territory coverage**, not data-center permits or
generation interconnection queues.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

from src.geo.fips import County, load_counties

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TERRITORY = ROOT / "data" / "raw" / "eia" / "Service_Territory_2024.xlsx"
DEFAULT_UTILITY = ROOT / "data" / "raw" / "eia" / "Utility_Data_2024.xlsx"

OUT_CROSSWALK = ROOT / "data" / "processed" / "crosswalks" / "eia_861_utility_county_v0.csv"
OUT_STATE = ROOT / "data" / "processed" / "panel" / "eia_861_utility_state_v0.csv"
OUT_SAMPLE = ROOT / "data" / "processed" / "panel" / "eia_861_utility_county_v0_sample.csv"
OUT_QA = ROOT / "data" / "processed" / "qa" / "eia_861_coverage_v0_qa.json"

TRANSFORM_VERSION = "eia_861_coverage_v0"
METRIC_LABEL = "utility_service_territory_coverage"
SAMPLE_N = 200

# Unique spelling / punctuation aliases → Census base name (same state).
# Only used when the alias resolves to exactly one gazetteer FIPS.
NAME_ALIASES: dict[tuple[str, str], str] = {
    ("LA", "desoto"): "de soto",
    ("LA", "la salle"): "lasalle",
    ("IL", "dewitt"): "de witt",
    ("TX", "dewitt"): "de witt",
    ("NM", "dona ana"): "dona ana",  # ñ stripped in norm
    ("FL", "miami dade"): "miami-dade",
    ("FL", "desoto"): "desoto",
    ("AK", "juneau"): "juneau",
    ("AK", "yakutat"): "yakutat",
    ("AK", "yukon koyukuk"): "yukon-koyukuk",
    ("AK", "prince of wales ketchikan"): "prince of wales-hyder",
}


def _ascii_fold(s: str) -> str:
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")


def _norm_name(s: str) -> str:
    s = _ascii_fold(str(s or "")).lower().strip()
    s = s.replace(".", "")
    s = s.replace("'", "")
    s = s.replace("saint ", "st ")
    s = re.sub(r"[\u2010-\u2015]", "-", s)
    s = re.sub(r"\s+", " ", s)
    return s


def _strip_county_suffix(name: str) -> str:
    n = name.strip()
    for suf in (
        " City and Borough",
        " Census Area",
        " Municipality",
        " Borough",
        " Parish",
        " County",
        " city",
    ):
        if n.endswith(suf):
            return n[: -len(suf)].strip()
    return n


def _build_county_index(
    counties: list[County],
) -> dict[tuple[str, str], list[County]]:
    idx: dict[tuple[str, str], list[County]] = defaultdict(list)
    for c in counties:
        base = _strip_county_suffix(c.name)
        for key_name in (base, c.name):
            idx[(c.state, _norm_name(key_name))].append(c)
    return idx


def _resolve_fips(
    state: str,
    county_raw: str,
    idx: dict[tuple[str, str], list[County]],
) -> tuple[Optional[str], str, Optional[str]]:
    """Return (county_fips, match_status, census_name).

    match_status: matched | ambiguous | unmatched | no_input
    """
    if not state or not county_raw:
        return None, "no_input", None
    if county_raw.strip().lower() in {"not applicable", "na", "n/a", "none"}:
        return None, "unmatched", None

    norm = _norm_name(county_raw)
    candidates = idx.get((state, norm), [])

    if not candidates and (state, norm) in NAME_ALIASES:
        candidates = idx.get((state, NAME_ALIASES[(state, norm)]), [])

    # Hyphen / space flexibility
    if not candidates:
        alt = norm.replace("-", " ")
        candidates = idx.get((state, alt), [])
    if not candidates:
        alt = norm.replace(" ", "-")
        candidates = idx.get((state, alt), [])

    uniq: dict[str, County] = {c.fips: c for c in candidates}
    if len(uniq) == 1:
        c = next(iter(uniq.values()))
        return c.fips, "matched", c.name
    if len(uniq) > 1:
        return None, "ambiguous", None
    return None, "unmatched", None


def _read_sheet_dicts(path: Path, sheet: str, header_row: int = 0) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet {sheet!r} not in {path.name}: {wb.sheetnames}")
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[header_row])]
    out: list[dict[str, Any]] = []
    for row in rows[header_row + 1 :]:
        if not row or all(v is None or v == "" for v in row):
            continue
        rec = {header[i]: row[i] if i < len(row) else None for i in range(len(header))}
        out.append(rec)
    return out


def _read_utility_states(path: Path) -> list[dict[str, Any]]:
    """Utility_Data has a title row then real header on row 1."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["States"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    # Find header containing 'Utility Number'
    header_idx = None
    for i, row in enumerate(rows[:5]):
        if row and "Utility Number" in [str(c) if c else "" for c in row]:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("Could not find Utility Number header in Utility_Data States sheet")
    header = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[header_idx])]
    out: list[dict[str, Any]] = []
    for row in rows[header_idx + 1 :]:
        if not row or row[1] is None:
            continue
        rec = {header[i]: row[i] if i < len(row) else None for i in range(len(header))}
        out.append(rec)
    return out


def transform(
    territory_path: Path = DEFAULT_TERRITORY,
    utility_path: Path = DEFAULT_UTILITY,
    *,
    crosswalk_csv: Path = OUT_CROSSWALK,
    state_csv: Path = OUT_STATE,
    sample_csv: Path = OUT_SAMPLE,
    qa_json: Path = OUT_QA,
    sample_n: int = SAMPLE_N,
) -> dict[str, Any]:
    if not territory_path.exists():
        raise FileNotFoundError(
            f"EIA Service Territory not found at {territory_path}. "
            "Extract from f8612024.zip or run: make fetch-eia"
        )
    if not utility_path.exists():
        raise FileNotFoundError(
            f"EIA Utility_Data not found at {utility_path}. Extract from f8612024.zip."
        )

    generated_at = datetime.now(timezone.utc)
    counties = load_counties()
    idx = _build_county_index(counties)
    known_fips = {c.fips for c in counties}

    territory_rows = _read_sheet_dicts(territory_path, "Counties_States")
    utility_rows = _read_utility_states(utility_path)

    status_counts: Counter[str] = Counter()
    by_state_match: dict[str, Counter[str]] = defaultdict(Counter)
    qa_exceptions: list[dict[str, Any]] = []
    crosswalk_out: list[dict[str, Any]] = []
    sample_rows: list[dict[str, Any]] = []

    n_territory = 0
    n_matched = 0
    n_ambiguous = 0
    n_unmatched = 0

    for rec in territory_rows:
        n_territory += 1
        state = str(rec.get("State") or "").strip().upper()
        county_raw = str(rec.get("County") or "").strip()
        util_num = rec.get("Utility Number")
        util_name = str(rec.get("Utility Name") or "").strip()
        data_year = rec.get("Data Year")

        fips, status, census_name = _resolve_fips(state, county_raw, idx)
        status_counts[status] += 1
        by_state_match[state or "?"][status] += 1

        if status == "matched":
            n_matched += 1
            assert fips is not None
            if fips not in known_fips:
                status = "unmatched"
                fips = None
                n_matched -= 1
                n_unmatched += 1
                status_counts["matched"] -= 1
                status_counts["unmatched"] += 1
        elif status == "ambiguous":
            n_ambiguous += 1
            if len(qa_exceptions) < 300:
                qa_exceptions.append(
                    {
                        "kind": "ambiguous_county_name",
                        "utility_number": util_num,
                        "utility_name": util_name,
                        "state": state,
                        "county_eia": county_raw,
                        "note": "County vs independent city share the same base name; FIPS not invented",
                    }
                )
        else:
            n_unmatched += 1
            if len(qa_exceptions) < 300:
                kind = "ct_legacy_county" if state == "CT" else "unmatched_county_name"
                qa_exceptions.append(
                    {
                        "kind": kind,
                        "utility_number": util_num,
                        "utility_name": util_name,
                        "state": state,
                        "county_eia": county_raw,
                        "note": (
                            "CT rows use pre-2022 county names; Census 2024 gazetteer uses COGs"
                            if state == "CT"
                            else "No unique gazetteer match; FIPS not invented"
                        ),
                    }
                )

        row = {
            "data_year": data_year if data_year is not None else "",
            "utility_number": util_num if util_num is not None else "",
            "utility_name": util_name,
            "state": state,
            "county_eia": county_raw,
            "county_fips": fips or "",
            "census_county_name": census_name or "",
            "match_status": status,
            "metric_label": METRIC_LABEL,
            "transform_version": TRANSFORM_VERSION,
            "grain_note": "annual_eia861_service_territory_not_monthly",
        }
        crosswalk_out.append(row)
        if len(sample_rows) < sample_n:
            sample_rows.append(row)

    # Utility → state stub (annual)
    state_out: list[dict[str, Any]] = []
    ownership: Counter[str] = Counter()
    for rec in utility_rows:
        state = str(rec.get("State") or "").strip().upper()
        util_num = rec.get("Utility Number")
        util_name = str(rec.get("Utility Name") or "").strip()
        own = str(rec.get("Ownership Type") or "").strip() or "unknown"
        nerc = str(rec.get("NERC Region") or "").strip()
        ownership[own] += 1
        state_out.append(
            {
                "data_year": rec.get("Data Year") if rec.get("Data Year") is not None else "",
                "utility_number": util_num if util_num is not None else "",
                "utility_name": util_name,
                "state": state,
                "ownership_type": own,
                "nerc_region": nerc,
                "metric_label": METRIC_LABEL,
                "transform_version": TRANSFORM_VERSION,
                "grain_note": "annual_eia861_utility_characteristics_state_not_county",
            }
        )

    crosswalk_fields = [
        "data_year",
        "utility_number",
        "utility_name",
        "state",
        "county_eia",
        "county_fips",
        "census_county_name",
        "match_status",
        "metric_label",
        "transform_version",
        "grain_note",
    ]
    crosswalk_csv.parent.mkdir(parents=True, exist_ok=True)
    with crosswalk_csv.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=crosswalk_fields)
        w.writeheader()
        for row in crosswalk_out:
            w.writerow(row)

    state_fields = [
        "data_year",
        "utility_number",
        "utility_name",
        "state",
        "ownership_type",
        "nerc_region",
        "metric_label",
        "transform_version",
        "grain_note",
    ]
    state_csv.parent.mkdir(parents=True, exist_ok=True)
    with state_csv.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=state_fields)
        w.writeheader()
        for row in state_out:
            w.writerow(row)

    with sample_csv.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=crosswalk_fields)
        w.writeheader()
        for row in sample_rows:
            w.writerow(row)

    exception_kinds = Counter(e["kind"] for e in qa_exceptions)
    qa: dict[str, Any] = {
        "transform_version": TRANSFORM_VERSION,
        "layer": "G",
        "metric_label": METRIC_LABEL,
        "content_scope": "eia861_utility_service_territory_and_utility_state",
        "explicitly_excluded": [
            "data_center_permit_counts",
            "monthly_panel_grain",
            "invented_county_fips",
        ],
        "raw_paths": {
            "service_territory": str(territory_path.relative_to(ROOT)),
            "utility_data": str(utility_path.relative_to(ROOT)),
            "zip_note": "Also present: data/raw/eia/f8612024.zip (source of XLSX extracts)",
        },
        "license_note": "US EIA / public domain",
        "geo_policy": (
            "Best-effort county_fips from EIA County name + State via Census gazetteer "
            "name match (suffix strip, St/Saint, hyphen/space, limited unique aliases). "
            "Ambiguous county-vs-city names and CT legacy counties left without FIPS."
        ),
        "grain": "annual (Data Year 2024); not county×month",
        "territory_rows": n_territory,
        "matched_fips": n_matched,
        "ambiguous": n_ambiguous,
        "unmatched": n_unmatched,
        "match_rate": round(n_matched / n_territory, 4) if n_territory else None,
        "match_status_counts": dict(status_counts),
        "utility_state_rows": len(state_out),
        "ownership_type_counts": dict(ownership.most_common()),
        "known_gaps": [
            {
                "gap": "connecticut_legacy_counties",
                "detail": (
                    "EIA Service Territory still lists CT legacy counties (Fairfield, "
                    "Hartford, …); Census 2024 gazetteer uses planning-region COGs. "
                    "No FIPS assigned without a dedicated CT crosswalk."
                ),
            },
            {
                "gap": "county_vs_independent_city",
                "detail": (
                    "VA/MD/MO names that collide with independent cities (e.g. Fairfax, "
                    "Baltimore, St Louis) are marked ambiguous — FIPS not invented."
                ),
            },
            {
                "gap": "alaska_historical_census_areas",
                "detail": (
                    "Some EIA AK labels (Valdez Cordova, Wrangell Petersburg, "
                    "Skagway Hoonah Angoon) reflect pre-split geographies; left unmatched."
                ),
            },
            {
                "gap": "not_monthly",
                "detail": "EIA-861 is annual; do not treat as county×month panel without interpolation policy.",
            },
        ],
        "qa_exception_kinds": dict(exception_kinds),
        "qa_exceptions_capped": qa_exceptions[:150],
        "outputs": {
            "utility_county_crosswalk_csv": str(crosswalk_csv.relative_to(ROOT)),
            "utility_state_csv": str(state_csv.relative_to(ROOT)),
            "sample_csv": str(sample_csv.relative_to(ROOT)),
            "qa_json": str(qa_json.relative_to(ROOT)),
        },
        "generated_at_utc": generated_at.isoformat(),
        "access_date_pt": "2026-09-16",
    }
    # Compact per-state match rates for inventory
    by_state_summary = {}
    for st, ctr in sorted(by_state_match.items()):
        tot = sum(ctr.values())
        by_state_summary[st] = {
            "n": tot,
            "matched": ctr.get("matched", 0),
            "ambiguous": ctr.get("ambiguous", 0),
            "unmatched": ctr.get("unmatched", 0) + ctr.get("no_input", 0),
            "match_rate": round(ctr.get("matched", 0) / tot, 4) if tot else None,
        }
    qa["by_state"] = by_state_summary

    qa_json.parent.mkdir(parents=True, exist_ok=True)
    qa_json.write_text(json.dumps(qa, indent=2) + "\n", encoding="utf-8")
    return qa


def main() -> None:
    p = argparse.ArgumentParser(
        description="EIA-861 → utility/state stub + best-effort utility→county FIPS (NOT DC permits)"
    )
    p.add_argument("--territory", type=Path, default=DEFAULT_TERRITORY)
    p.add_argument("--utility", type=Path, default=DEFAULT_UTILITY)
    p.add_argument("--sample-n", type=int, default=SAMPLE_N)
    args = p.parse_args()
    qa = transform(
        territory_path=args.territory,
        utility_path=args.utility,
        sample_n=args.sample_n,
    )
    print(f"transform_version={qa['transform_version']}")
    print(f"metric_label={qa['metric_label']}")
    print(f"territory_rows={qa['territory_rows']}")
    print(f"matched_fips={qa['matched_fips']}")
    print(f"ambiguous={qa['ambiguous']}")
    print(f"unmatched={qa['unmatched']}")
    print(f"match_rate={qa['match_rate']}")
    print(f"utility_state_rows={qa['utility_state_rows']}")
    print(f"outputs={qa['outputs']}")


if __name__ == "__main__":
    main()
